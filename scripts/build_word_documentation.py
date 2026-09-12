#!/usr/bin/env python3
"""
Build Beautiful Word (.docx) Documents for VaxPlan:
1. VaxPlan Business Case & Stakeholder Investment Prospectus
2. VaxPlan Organizational Terms of Reference (ToR) & Operational Manual
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

DOCS_DIR = os.path.join(os.path.dirname(__file__), "..", "docs")

# Visual Palette
COLOR_PRIMARY_NAVY = "0F172A"   # Slate 900
COLOR_DEEP_BLUE    = "1E3A8A"   # Blue 900
COLOR_ACCENT_BLUE  = "2563EB"   # Blue 600
COLOR_EMERALD      = "059669"   # Emerald 600
COLOR_TEXT_MAIN    = "1E293B"   # Slate 800
COLOR_TEXT_MUTED   = "64748B"   # Slate 500
COLOR_ZEBRA        = "F8FAFC"   # Slate 50
COLOR_BORDER       = "CBD5E1"   # Slate 300
COLOR_CARD_BG      = "F1F5F9"   # Slate 100
COLOR_CALLOUT_BG   = "EFF6FF"   # Blue 50
COLOR_CALLOUT_BORDER = "3B82F6" # Blue 500

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''<w:tcMar {nsdecls("w")}>
        <w:top w:w="{top}" w:type="dxa"/>
        <w:bottom w:w="{bottom}" w:type="dxa"/>
        <w:left w:w="{left}" w:type="dxa"/>
        <w:right w:w="{right}" w:type="dxa"/>
    </w:tcMar>''')
    tcPr.append(tcMar)

def set_cell_borders(cell, top="none", bottom="none", left="none", right="none", 
                     color="CBD5E1", sz="4"):
    tcPr = cell._element.get_or_add_tcPr()
    borders_elm = parse_xml(f'''<w:tcBorders {nsdecls("w")}>
        <w:top w:val="{top}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:left w:val="{left}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:bottom w:val="{bottom}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:right w:val="{right}" w:sz="{sz}" w:space="0" w:color="{color}"/>
    </w:tcBorders>''')
    tcPr.append(borders_elm)

def style_table_header(row, col_names, bg_hex=COLOR_PRIMARY_NAVY, text_hex="FFFFFF"):
    for i, name in enumerate(col_names):
        cell = row.cells[i]
        set_cell_background(cell, bg_hex)
        set_cell_margins(cell, top=140, bottom=140, left=150, right=150)
        set_cell_borders(cell, top="single", bottom="single", left="none", right="none", color=COLOR_DEEP_BLUE, sz="8")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(name)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor.from_string(text_hex)

def format_data_cell(cell, text, align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, font_size=9.5, 
                      color=COLOR_TEXT_MAIN, bg_color=None):
    if bg_color:
        set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=90, bottom=90, left=130, right=130)
    set_cell_borders(cell, top="single", bottom="single", left="none", right="none", color=COLOR_BORDER, sz="4")
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(font_size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

def add_callout_box(doc, text, bold_prefix="", border_color=COLOR_CALLOUT_BORDER, bg_color=COLOR_CALLOUT_BG):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.rows[0].cells[0]
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    set_cell_borders(cell, top="none", bottom="none", left="single", right="none", color=border_color, sz="24") # 3pt left border
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    if bold_prefix:
        r_prefix = p.add_run(bold_prefix + " ")
        r_prefix.bold = True
        r_prefix.font.name = "Calibri"
        r_prefix.font.size = Pt(10)
        r_prefix.font.color.rgb = RGBColor.from_string(border_color)
    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(9.5)
    r_text.font.italic = True
    r_text.font.color.rgb = RGBColor.from_string(COLOR_TEXT_MAIN)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_code_diagram_box(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.rows[0].cells[0]
    set_cell_background(cell, COLOR_CARD_BG)
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color=COLOR_BORDER, sz="6")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = Pt(12)
    for i, line in enumerate(lines):
        run = p.add_run(line + ("\n" if i < len(lines) - 1 else ""))
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string(COLOR_DEEP_BLUE)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(15)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(COLOR_PRIMARY_NAVY)
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(12.5)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(COLOR_DEEP_BLUE)
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(COLOR_TEXT_MAIN)
    return p

def add_p(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(COLOR_TEXT_MAIN)
    return p

def add_bullet(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + ": ")
        r_pre.bold = True
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(9.5)
        r_pre.font.color.rgb = RGBColor.from_string(COLOR_DEEP_BLUE)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(COLOR_TEXT_MAIN)
    return p

def setup_headers_footers(doc, title_short):
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.different_first_page_header_footer = True

    # Header for body pages
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run(f"VaxPlan Platform • {title_short}")
    hrun.font.name = "Calibri"
    hrun.font.size = Pt(8.5)
    hrun.font.color.rgb = RGBColor.from_string(COLOR_TEXT_MUTED)

    # Footer for body pages
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = fp.add_run("VaxPlan © 2026 • Confidential & Strategic Health Investment Document")
    frun.font.name = "Calibri"
    frun.font.size = Pt(8)
    frun.font.color.rgb = RGBColor.from_string(COLOR_TEXT_MUTED)

def add_cover_page(doc, title, subtitle, meta_items):
    # Top spacing
    p_top = doc.add_paragraph()
    p_top.paragraph_format.space_before = Pt(36)
    p_top.paragraph_format.space_after = Pt(12)
    
    # Organization / Platform Tag
    p_org = doc.add_paragraph()
    p_org.paragraph_format.space_after = Pt(6)
    r_org = p_org.add_run("VAXPLAN PLATFORM • GLOBAL HEALTH INITIATIVE")
    r_org.font.name = "Calibri"
    r_org.font.size = Pt(9.5)
    r_org.bold = True
    r_org.font.color.rgb = RGBColor.from_string(COLOR_ACCENT_BLUE)

    # Main Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(10)
    p_title.paragraph_format.line_spacing = 1.1
    r_title = p_title.add_run(title)
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(26)
    r_title.bold = True
    r_title.font.color.rgb = RGBColor.from_string(COLOR_PRIMARY_NAVY)

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(24)
    r_sub = p_sub.add_run(subtitle)
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(12.5)
    r_sub.font.color.rgb = RGBColor.from_string(COLOR_TEXT_MUTED)

    # Decorative Line
    p_line = doc.add_paragraph()
    p_line.paragraph_format.space_after = Pt(28)
    r_line = p_line.add_run("―" * 42)
    r_line.font.color.rgb = RGBColor.from_string(COLOR_ACCENT_BLUE)
    r_line.bold = True

    # Metadata Card Table
    table = doc.add_table(rows=len(meta_items), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(4.3)
    
    for i, (k, v) in enumerate(meta_items):
        row = table.rows[i]
        bg = COLOR_CARD_BG if i % 2 == 0 else "FFFFFF"
        format_data_cell(row.cells[0], k, align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, font_size=9, color=COLOR_DEEP_BLUE, bg_color=bg)
        format_data_cell(row.cells[1], v, align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, font_size=9, color=COLOR_TEXT_MAIN, bg_color=bg)

    p_bot = doc.add_paragraph()
    p_bot.paragraph_format.space_before = Pt(40)
    r_bot = p_bot.add_run("Reach every child. Plan every session. Optimize every dose.")
    r_bot.font.name = "Calibri"
    r_bot.font.size = Pt(10)
    r_bot.font.italic = True
    r_bot.font.color.rgb = RGBColor.from_string(COLOR_EMERALD)

    doc.add_page_break()

# ==============================================================================
# 1. BUILD BUSINESS CASE DOCUMENT
# ==============================================================================
def build_business_case_docx():
    doc = docx.Document()
    setup_headers_footers(doc, "Business Case & Investment Prospectus")

    add_cover_page(
        doc,
        title="VaxPlan: Enterprise Business Case & Stakeholder Investment Prospectus",
        subtitle="Digital GIS Microplanning, Dual-Source Populations & WHO RED-Aligned Operational Excellence for Universal Immunization",
        meta_items=[
            ("Document Version", "v1.0 (Comprehensive Investment Edition)"),
            ("Effective Date", "September 2026"),
            ("Target Audience", "Ministries of Health, National EPI Managers, Gavi, WHO, UNICEF, BMGF, Global Fund"),
            ("Global Health Alignment", "WHO RED Strategy (Tools 1-10), WHO IIP Module 4, Gavi 5.0, IA2030"),
            ("Platform State", "Production-Ready (Web, Offline PWA, Android, Windows Desktop)"),
            ("Sovereign Hosting", "Supports MoH Sovereign National Cloud & Hybrid Multi-Tenant Deployment"),
        ]
    )

    add_h1(doc, "1. Executive Summary & Investment Thesis")
    add_p(doc, "Immunization remains one of humanity's most cost-effective public health interventions, saving millions of lives annually. Yet globally, over 14 million children each year are 'zero-dose'—failing to receive even a single dose of diphtheria, tetanus, and pertussis (DTP-1) vaccine—and an additional 6 million are severely under-immunized.")
    add_p(doc, "The primary bottleneck to universal coverage is no longer vaccine manufacturing capacity or national cold stores. It is the systemic breakdown of operational planning at the last mile. In low- and middle-income countries (LMICs), national microplans overwhelmingly exist as disconnected paper forms, static spreadsheets, and hand-drawn wall sketch maps. Clinic managers calculate targets with outdated census projections, plan outreach sessions without terrain intelligence, and discover stockouts or canceled visits weeks after children have been turned away.")

    add_callout_box(
        doc,
        "VaxPlan is a mobile-first, offline-capable GIS microplanning platform that takes a Ministry of Health from an interactive satellite map of its facilities all the way down to a vaccinator's daily itinerary—and back up to real-time national equity dashboards. VaxPlan operationalizes the WHO Reach Every District (RED) strategy into a modern digital system.",
        bold_prefix="The VaxPlan Vision:"
    )

    add_h2(doc, "The Continuous Operational Cycle")
    add_code_diagram_box(doc, [
        "  [1. SPATIAL MAP]      -->  [2. RED-ALIGNED WIZARD]  -->  [3. LOGISTICS & TEAMS]",
        "  High-Res GIS,               Prepopulated Worksheets,      Session Sizing, Cold",
        "  Boundaries, Catchments      Dual-Source Denominators      Chain, Fuel & Budgets",
        "          ^                                                           |",
        "          |                                                           v",
        "  [6. GOVERNANCE]       <--  [5. COVERAGE & DEFAULTERS] <--  [4. FIELD EXECUTION]",
        "  7-Day Review Window,        QR Health Cards, Client       Offline PWA Mobile,",
        "  Approval Audit Trail        Register, Traceability        Daily Itineraries"
    ])

    add_h1(doc, "2. The Problem: The Anatomy of Last-Mile Planning Failure")
    add_p(doc, "Immunization programs across Sub-Saharan Africa, South Asia, and the Pacific encounter five structural bottlenecks that waste funding and leave communities unreached:")

    add_bullet(doc, "Health facilities rely on national censuses that are 10–20 years outdated, completely missing rapid urbanization, displaced settlements, and nomadic migration corridors. Denominators become arbitrary guesswork.", bold_prefix="The Denominator Dilemma")
    add_bullet(doc, "Outreach sessions scheduled on paper are canceled in practice due to unbudgeted transport, seasonal river flooding, or impassable roads. Supervisors remain unaware until quarterly reports months later.", bold_prefix="The Ghost Session Phenomenon")
    add_bullet(doc, "Adjoining health centers lack digitized boundaries. Overlapping catchments lead to friction, while border settlements fall into dead zones where each facility assumes the other is responsible.", bold_prefix="Contested Catchment Zones")
    add_bullet(doc, "Microplans operate disconnected from physical stock ledgers. Sessions are planned for remote sites with vaccine carriers lacking certified ice packs, leading to heat spoilage or freeze damage.", bold_prefix="Cold Chain Disconnect")
    add_bullet(doc, "Health staff spend dozens of hours retyping identical demographic figures and antigen counts across redundant ledgers and donor reporting forms.", bold_prefix="Heavy Administrative Overhead")

    add_h1(doc, "3. Platform Differentiators & Comparative Advantage")
    
    table_diff = doc.add_table(rows=8, cols=3)
    table_diff.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_diff.autofit = False
    table_diff.columns[0].width = Inches(1.8)
    table_diff.columns[1].width = Inches(2.2)
    table_diff.columns[2].width = Inches(2.5)

    headers = ["Functional Area", "Legacy Spreadsheets & Paper", "The VaxPlan Platform Advantage"]
    style_table_header(table_diff.rows[0], headers)

    data_diff = [
        ("Microplanning", "Complex paper forms, unvalidated Excel files", "Guided 12-step wizard directly mapped to WHO RED 10-step worksheets with auto-prepopulation."),
        ("Catchment Mapping", "Static wall sketch maps, no coordinates", "Interactive PostGIS spatial engine, GeoJSON polygons, overlap detection, and conflict resolution."),
        ("Population Targets", "Single outdated census projection", "Dual-source engine: Official NSO census reconciled with WorldPop 100m gridded satellite population."),
        ("Field Connectivity", "Online-only tools that fail in remote bush", "True offline-first PWA with IndexedDB local storage and deterministic sync reconciliation."),
        ("Plan Approval & Review", "Informal verbal approval, unmonitored", "Multi-tier approval workflow with mandatory 7-day review windows and configurable implementation lead times."),
        ("Version Integrity", "Files overwritten, lost history", "Immutable snapshot versioning: automatic version created on open, edit, save, and close."),
        ("Supply Forecasting", "Disjointed manual guesses", "Automated antigen, AD syringe, dilution syringe, safety box, and cold box sizing based on WHO vial wastage."),
    ]

    for idx, (f_area, leg, adv) in enumerate(data_diff, start=1):
        row = table_diff.rows[idx]
        bg = COLOR_ZEBRA if idx % 2 == 1 else "FFFFFF"
        format_data_cell(row.cells[0], f_area, bold=True, bg_color=bg)
        format_data_cell(row.cells[1], leg, color=COLOR_TEXT_MUTED, bg_color=bg)
        format_data_cell(row.cells[2], adv, color=COLOR_DEEP_BLUE, bold=True, bg_color=bg)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_h1(doc, "4. Quantified Business Case & Return on Investment (ROI)")
    add_p(doc, "VaxPlan delivers measurable economic and operational efficiencies for national governments and donor sponsors:")

    add_bullet(doc, "Elimination of redundant vehicle journeys and canceled outreach sessions reduces per-session transport expenditure by 18% to 25% through spatial route clustering.", bold_prefix="18–25% Transport Savings")
    add_bullet(doc, "Precise quarterly forecasting by antigen presentation reduces vaccine expiration and open-vial wastage by 15% to 30%, protecting expensive pentavalent and PCV investments.", bold_prefix="15–30% Wastage Reduction")
    add_bullet(doc, "Automated data prepopulation returns an estimated 120,000 health worker hours per national planning cycle from paperwork to frontline clinical care.", bold_prefix="120,000+ Health Worker Hours Saved")
    add_bullet(doc, "High-resolution satellite population overlays identify unreached and zero-dose communities, directing operational budgets precisely where outbreak risk is concentrated.", bold_prefix="Zero-Dose Targeting Precision")

    add_callout_box(
        doc,
        "Representative Country Scenario (1,500 health facilities, 400,000 annual target birth cohort): Direct annual operational savings of $420,000 - $770,000 in transport, fuel, and preserved vaccine stock, achieving complete investment recovery within 9 to 14 months of nationwide deployment.",
        bold_prefix="Economic Payback Horizon:"
    )

    add_h1(doc, "5. Governance, Multi-Tier Approvals & Policy Enforcement")
    add_p(doc, "Microplan governance in VaxPlan is enforced through rigorous, policy-driven workflows:")
    add_bullet(doc, "Microplans cannot be approved immediately upon submission. A mandatory review window (default 7 days, configurable per country tenant) is enforced, providing district supervisors with the requisite desk-review and audit period.", bold_prefix="7-Day Review Window")
    add_bullet(doc, "Vaccination sessions must be scheduled for execution at least X days (configurable, default 7 days) following approval, guaranteeing that vaccines are distributed and community mobilizers have sufficient notice.", bold_prefix="Implementation Lead-Time Policy")
    add_bullet(doc, "Each microplan captures a distinct point-in-time snapshot whenever opened, edited, saved, submitted, approved, or restored, creating an unalterable audit trail for sovereign health authorities and donors.", bold_prefix="Full Draft Lifecycle Versioning")

    add_h1(doc, "6. Commercial Models & Sovereign Deployment Frameworks")
    add_p(doc, "VaxPlan offers flexible, sustainable engagement models designed for national ownership:")
    add_bullet(doc, "Annual or multi-year national agreement covering sovereign cloud hosting, system maintenance, 24/7 technical support, and continuous feature updates.", bold_prefix="1. Sovereign MoH Enterprise License")
    add_bullet(doc, "Multilateral donor funding (Gavi Operational Costs, UNICEF Health Systems Strengthening, BMGF) covering initial baseline GIS mapping, training, and 24-month transition to full MoH management.", bold_prefix="2. Donor-Sponsored Technical Assistance")
    add_bullet(doc, "Rapid implementation in priority zero-dose districts, border areas, or humanitarian refugee corridors in partnership with international health NGOs.", bold_prefix="3. Targeted Sub-National Corridor Pilots")

    # Save Business Case document
    out_path = os.path.join(DOCS_DIR, "VAXPLAN_BUSINESS_CASE_AND_STAKEHOLDER_PROSPECTUS.docx")
    doc.save(out_path)
    print(f"Successfully generated: {out_path}")

# ==============================================================================
# 2. BUILD ORGANIZATIONAL TOR DOCUMENT
# ==============================================================================
def build_organizational_tor_docx():
    doc = docx.Document()
    setup_headers_footers(doc, "Organizational ToR & Operating Manual")

    add_cover_page(
        doc,
        title="VaxPlan: Organizational Terms of Reference (ToR) & Operational Manual",
        subtitle="Institutional Governance Charters, Executive Mandates, Consultant Scopes of Work, and Team Standard Operating Procedures",
        meta_items=[
            ("Document Version", "v1.0 (Governance & Operations Edition)"),
            ("Effective Date", "September 2026"),
            ("Target Audience", "Founder, Co-Founders, Advisory Board, Consultants, Core Employees, MoH Liaisons"),
            ("Governance Framework", "Sovereign Health Data Sovereignty, Zero-Data-Loss Standards, Role-Based Access Control"),
            ("Standard Alignment", "WHO Digital Health Guidelines, ISO/IEC 27001 Security Controls, HL7 FHIR Interoperability"),
            ("Review Frequency", "Annual Executive Review or upon Country Deployment Milestone"),
        ]
    )

    add_h1(doc, "1. Organizational Mandate & Core Values")
    add_p(doc, "The mission of VaxPlan is to eradicate preventable childhood disease and reach every zero-dose community by providing Ministries of Health, frontline health workers, and global partners with an intelligent, offline-first GIS microplanning and operations platform.")
    
    add_callout_box(
        doc,
        "Every line of code, technical architecture decision, and in-country training initiative must prioritize the frontline health worker. We operate under strict principles: Health Worker First, Offline-Resilient by Default, Absolute Data Sovereignty, Non-Destructive System Integrity, and WHO Evidence-Based Standards.",
        bold_prefix="Operational Doctrine:"
    )

    add_h1(doc, "2. Executive Leadership Charters (ToR)")

    add_h2(doc, "2.1 Founder & Chief Executive Officer (CEO)")
    add_bullet(doc, "Lead global vision and multi-year organizational growth to establish VaxPlan as the recognized national microplanning standard across 20+ countries.", bold_prefix="Strategic Vision")
    add_bullet(doc, "Lead sovereign ministerial diplomacy with Ministers of Health, National EPI Directors, and donor executives at Gavi, WHO, UNICEF, and the Gates Foundation.", bold_prefix="Ministerial Partnerships")
    add_bullet(doc, "Secure institutional financing, technical assistance grants, and sustainable multi-year software agreements.", bold_prefix="Capitalization")
    add_bullet(doc, "Lead executive performance, uphold organizational ethics, and champion global health equity.", bold_prefix="Governance")

    add_h2(doc, "2.2 Co-Founder & Chief Technology Officer (CTO)")
    add_bullet(doc, "Architect and govern the full technology stack: React frontend, Node.js backend, PostGIS spatial database, and offline-first IndexedDB replication.", bold_prefix="Platform Architecture")
    add_bullet(doc, "Guarantee sub-second GIS mapping performance, polygon buffering, conflict-free offline replication, and fault-tolerant background data sync.", bold_prefix="Spatial & Offline Leadership")
    add_bullet(doc, "Enforce strict multi-tenant row-level security, tenant data isolation, TLS 1.3 encryption, and zero-data-loss additive database practices.", bold_prefix="Security & Data Safety")
    add_bullet(doc, "Mentor senior engineers, oversee CI/CD automation, and ensure 100% test passing rates across all release builds.", bold_prefix="Engineering Excellence")

    add_h2(doc, "2.3 Co-Founder / Head of Public Health & Immunization Strategy")
    add_bullet(doc, "Audit every workflow against WHO Reach Every District (RED Tools 1-10), WHO Immunization in Practice (IIP Module 4), and Gavi 5.0 benchmarks.", bold_prefix="WHO RED & EPI Compliance")
    add_bullet(doc, "Formulate mathematical algorithms for population disaggregation, dropout calculations, antigen wastage rates, and cold-chain capacity sizing.", bold_prefix="Epidemiological Modeling")
    add_bullet(doc, "Lead technical dialogues with National EPI Technical Working Groups (TWGs) and WHO/UNICEF country office epidemiologists.", bold_prefix="Stakeholder Advisory")
    add_bullet(doc, "Design rigorous impact measurement protocols to evaluate reductions in zero-dose children and missed opportunities for vaccination (MOV).", bold_prefix="Impact Evaluation")

    add_h2(doc, "2.4 Chief Operating Officer (COO) / Head of Deployments")
    add_bullet(doc, "Develop and manage country onboarding Gantt charts, tracking milestones from initial GIS data gathering to nationwide sign-off.", bold_prefix="Deployment Roadmaps")
    add_bullet(doc, "Coordinate hardware deployment (tablets, rugged devices, solar battery packs) and local connectivity support for district field teams.", bold_prefix="Field Readiness")
    add_bullet(doc, "Oversee country deployment operational budgets, consultant field logistics, and compliance with donor accounting guidelines.", bold_prefix="Budget Oversight")
    add_bullet(doc, "Ensure tight cross-functional synergy between software engineers, public health advisors, and in-country field deployment teams.", bold_prefix="Operational Coordination")

    add_h1(doc, "3. Specialized Consultants Terms of Reference (ToR)")
    
    table_cons = doc.add_table(rows=5, cols=3)
    table_cons.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_cons.autofit = False
    table_cons.columns[0].width = Inches(1.8)
    table_cons.columns[1].width = Inches(2.2)
    table_cons.columns[2].width = Inches(2.5)

    headers_c = ["Consultant Role", "Core Scope of Work", "Key Deliverables"]
    style_table_header(table_cons.rows[0], headers_c)

    data_cons = [
        ("Senior Immunization Policy Consultant", 
         "Review national EPI schedules, local antigen policies, and disease calendars to configure country tenants.",
         "Country Immunization Matrix, Readiness Assessment Report, and Ministerial Policy Brief."),
        ("GIS & Spatial Analytics Specialist", 
         "Harmonize administrative boundary Shapefiles (Admin 0-4), village coordinates, and 100m WorldPop rasters.",
         "Harmonized National Spatial GeoJSON Registry, Overlap Conflict Audit, Printable Catchment Maps."),
        ("Health Systems Interoperability Specialist", 
         "Architect bi-directional sync between VaxPlan and national DHIS2 instances and electronic immunization registries.",
         "OpenHIE/FHIR Data Exchange Specification, Validated DHIS2 Export Adapter, Integration Test Report."),
        ("Health Economics & M&E Specialist", 
         "Evaluate logistics cost reductions, transport optimization, and vaccine wastage prevention across deployment districts.",
         "Cost-Effectiveness Evaluation Report, Zero-Dose Reduction Impact Study, and Donor Grant Audit File."),
    ]

    for idx, (role, scope, deliv) in enumerate(data_cons, start=1):
        row = table_cons.rows[idx]
        bg = COLOR_ZEBRA if idx % 2 == 1 else "FFFFFF"
        format_data_cell(row.cells[0], role, bold=True, bg_color=bg)
        format_data_cell(row.cells[1], scope, color=COLOR_TEXT_MAIN, bg_color=bg)
        format_data_cell(row.cells[2], deliv, color=COLOR_DEEP_BLUE, bold=True, bg_color=bg)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_h1(doc, "4. Core Employee Mandates (ToR)")
    add_bullet(doc, "Lead backend and API engineering in TypeScript/Node.js; optimize PostGIS spatial queries; maintain 90%+ automated test coverage; ensure backwards-compatible schema evolution.", bold_prefix="Lead Full-Stack Engineer")
    add_bullet(doc, "Build high-performance mapping interfaces (Leaflet/MapLibre); optimize offline IndexedDB caching; ensure 60fps responsiveness on low-cost Android hardware.", bold_prefix="GIS & Offline Frontend Engineer")
    add_bullet(doc, "Conduct hands-on training for district health teams and nurses; troubleshoot microplanning cycles; capture user friction points and translate into engineering tickets.", bold_prefix="Field Implementation & Training Lead")
    add_bullet(doc, "Execute end-to-end regression test suites; validate mathematical forecasting across all currency and cohort variations; conduct offline disconnect/reconnect stress tests.", bold_prefix="Quality Assurance & Data Integrity Engineer")

    add_h1(doc, "5. Standard Operating Procedures (SOPs) for Team Operations")

    add_h2(doc, "5.1 Non-Destructive Deployment & Zero-Data-Loss SOP")
    add_bullet(doc, "Production environment files (.env), Nginx server configurations, and process manager settings (PM2) are protected assets and must NEVER be overwritten by deployment scripts.", bold_prefix="Protected Configurations")
    add_bullet(doc, "Destructive SQL commands (DROP TABLE, TRUNCATE, destructive column drops) are strictly forbidden in production. All schema evolution must be additive, nullable, and backward-compatible.", bold_prefix="Additive Schema Evolution")
    add_bullet(doc, "Every microplan must generate an immutable version record whenever opened, edited, saved, submitted, or approved.", bold_prefix="Immutable Snapshot Versioning")
    add_bullet(doc, "Every software release must achieve 0 TypeScript errors (npm run check), 100% test pass rate (npm test), and a clean production build (npm run build) prior to release.", bold_prefix="Pre-Release Gates")

    add_h2(doc, "5.2 Microplan Governance & Quality Assurance SOP")
    add_bullet(doc, "A minimum review window (default 7 days) is enforced between plan submission and approval to guarantee thorough desk reviews by district supervisors.", bold_prefix="7-Day Review Window")
    add_bullet(doc, "All vaccination sessions must be scheduled to occur at least X days (default 7 days) after plan approval, ensuring vaccines are requisitioned and distributed.", bold_prefix="Session Implementation Lead-Time")
    add_bullet(doc, "Supervisors must use the standardized WHO RED supervision checklist during quarterly supportive supervision visits, with findings tied to action items.", bold_prefix="Supervision Action Register")

    add_h1(doc, "6. Code of Conduct, Ethics & Data Privacy")
    add_p(doc, "All employees, contractors, and consultants must adhere to strict ethical standards:")
    add_bullet(doc, "Individual patient identities and child records are strictly protected health data and must never be shared or commercialized.", bold_prefix="Patient Anonymity")
    add_bullet(doc, "Team members must disclose any commercial or advisory interests with pharmaceutical or medical equipment manufacturers.", bold_prefix="Conflict of Interest")
    add_bullet(doc, "We view community health volunteers and nurses as co-creators of VaxPlan. We listen with humility and act rapidly on frontline feedback.", bold_prefix="Frontline Respect")

    # Save Organizational ToR document
    out_path = os.path.join(DOCS_DIR, "VAXPLAN_ORGANIZATIONAL_TOR_AND_OPERATING_MANUAL.docx")
    doc.save(out_path)
    print(f"Successfully generated: {out_path}")

if __name__ == "__main__":
    build_business_case_docx()
    build_organizational_tor_docx()
